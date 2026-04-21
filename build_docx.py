"""Generate NeuroSentinel AI Guide DOCX from two markdown parts."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re

ARTIFACTS = r'C:\Users\manik\.gemini\antigravity\brain\3be35c0b-8ecb-4843-b406-8b999313aadd\artifacts'
p1 = open(f'{ARTIFACTS}/NeuroSentinel_Guide_April19_Part1.md', encoding='utf-8').read()
p2 = open(f'{ARTIFACTS}/NeuroSentinel_Guide_April19_Part2.md', encoding='utf-8').read()
content = p1 + '\n\n' + p2

doc = Document()
for section in doc.sections:
    section.top_margin = Cm(2.5); section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.8); section.right_margin = Cm(2.8)

# ── Title page ──────────────────────────────────────────────────────────────
t = doc.add_heading('NeuroSentinel AI', 0)
t.runs[0].font.size = Pt(32)
t.runs[0].font.color.rgb = RGBColor(0x0D, 0x47, 0xA1)
t.alignment = WD_ALIGN_PARAGRAPH.CENTER

sub = doc.add_paragraph('Complete Project & Research Guide')
sub.runs[0].font.size = Pt(16); sub.runs[0].bold = True
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

d = doc.add_paragraph('Version 2.1  |  April 19, 2026')
d.runs[0].font.size = Pt(11); d.alignment = WD_ALIGN_PARAGRAPH.CENTER

d2 = doc.add_paragraph('SCOUT \u2014 Seizure Clinical Operations & Understanding Tool')
d2.runs[0].font.size = Pt(10); d2.runs[0].italic = True
d2.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_page_break()


def blue_cell(cell):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), '1565C0'); tcPr.append(shd)


def stripe_cell(cell):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'E3F2FD'); tcPr.append(shd)


def code_bg(para):
    pPr = para._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F1F8E9'); pPr.append(shd)


def add_table(doc, lines):
    rows = [l for l in lines if l.startswith('|')]
    if len(rows) < 2:
        return
    header = [c.strip() for c in rows[0].strip('|').split('|')]
    data = [
        [c.strip() for c in r.strip('|').split('|')]
        for r in rows[2:]
        if not all(c.strip().replace('-', '') == '' for c in r.strip('|').split('|'))
    ]
    nc = len(header)
    tbl = doc.add_table(rows=1 + len(data), cols=nc)
    tbl.style = 'Table Grid'
    hr = tbl.rows[0]
    for i, h in enumerate(header[:nc]):
        cell = hr.cells[i]; cell.text = h
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True; run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(255, 255, 255)
        blue_cell(cell)
    for ri, dr in enumerate(data):
        row = tbl.rows[ri + 1]
        for ci, val in enumerate(dr[:nc]):
            row.cells[ci].text = val
            for para in row.cells[ci].paragraphs:
                for run in para.runs:
                    run.font.size = Pt(9)
        if ri % 2 == 1:
            for ci in range(nc):
                stripe_cell(row.cells[ci])
    doc.add_paragraph()


# ── Parse markdown ──────────────────────────────────────────────────────────
BACKTICK3 = '```'

in_code = False
code_lines = []
table_lines = []
lines = content.split('\n')
i = 0
while i < len(lines):
    line = lines[i]

    # Code fence detection
    stripped = line.strip()
    if stripped.startswith(BACKTICK3):
        if not in_code:
            in_code = True; code_lines = []
        else:
            in_code = False
            if code_lines:
                para = doc.add_paragraph()
                para.style = 'No Spacing'
                run = para.add_run('\n'.join(code_lines))
                run.font.name = 'Courier New'
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor(0x1B, 0x5E, 0x20)
                code_bg(para)
                doc.add_paragraph()
            code_lines = []
        i += 1; continue

    if in_code:
        code_lines.append(line); i += 1; continue

    # Tables
    if line.startswith('|'):
        table_lines.append(line); i += 1; continue
    else:
        if table_lines:
            add_table(doc, table_lines)
            table_lines = []

    # Headings
    if line.startswith('#### '):
        h = doc.add_heading(line[5:], 4)
        for r in h.runs: r.font.size = Pt(11)
    elif line.startswith('### '):
        h = doc.add_heading(line[4:], 3)
        for r in h.runs:
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(0x1A, 0x23, 0x7E)
    elif line.startswith('## '):
        h = doc.add_heading(line[3:], 2)
        for r in h.runs:
            r.font.size = Pt(14)
            r.font.color.rgb = RGBColor(0x0D, 0x47, 0xA1)
    elif line.startswith('# ') and 'NeuroSentinel AI \u2014' not in line:
        h = doc.add_heading(line[2:], 1)
        for r in h.runs:
            r.font.size = Pt(16)
            r.font.color.rgb = RGBColor(0x0D, 0x47, 0xA1)
    elif line.startswith('---'):
        doc.add_paragraph()
    elif line.strip() == '':
        pass
    elif line.startswith('- ') or line.startswith('* '):
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(line[2:]).font.size = Pt(10)
    elif re.match(r'^\d+\. ', line):
        p = doc.add_paragraph(style='List Number')
        p.add_run(re.sub(r'^\d+\. ', '', line)).font.size = Pt(10)
    elif line.startswith('> '):
        p = doc.add_paragraph(line[2:])
        if p.runs:
            p.runs[0].italic = True
            p.runs[0].font.size = Pt(10)
            p.runs[0].font.color.rgb = RGBColor(0x37, 0x47, 0x6F)
    else:
        p = doc.add_paragraph()
        for part in re.split(r'(\*\*.*?\*\*)', line):
            if part.startswith('**') and part.endswith('**'):
                r = p.add_run(part[2:-2]); r.bold = True; r.font.size = Pt(10)
            else:
                r = p.add_run(part); r.font.size = Pt(10)
    i += 1

if table_lines:
    add_table(doc, table_lines)

OUT = (r'c:\Users\manik\OneDrive\Desktop\Adwith\Projects\NeuroSentinel AI'
       r'\NeuroSentinel AI\seizure\NeuroSentinel_AI_Guide_April19_2026.docx')
doc.save(OUT)
print('Saved:', OUT)
